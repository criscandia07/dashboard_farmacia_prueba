import io
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def export_schedule_to_excel(df_schedule, params, funciones, kpis):
    """Genera un archivo Excel estructurado en 3 hojas con colores y estilos."""
    wb = openpyxl.Workbook()
    
    # ---------------------------------------------------------
    # HOJA 1: PLANILLA
    # ---------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Planilla de Turnos"
    
    # Estilos Base
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    color_map = {f["nombre"]: f["color"].replace("#", "") for f in funciones}
    
    # Encabezados
    headers = list(df_schedule.columns)
    ws1.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws1.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Filas de datos
    for r_idx, row in df_schedule.iterrows():
        row_data = list(row)
        ws1.append(row_data)
        curr_row = r_idx + 2
        for c_idx, val in enumerate(row_data, 1):
            cell = ws1.cell(row=curr_row, column=c_idx)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            # Pintar función con su color pastel
            if str(val) in color_map:
                hex_c = color_map[str(val)]
                cell.fill = PatternFill(start_color=hex_c, end_color=hex_c, fill_type="solid")
                cell.font = Font(name="Segoe UI", size=10, bold=True)

    # Autoajuste de columnas
    for col in ws1.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # ---------------------------------------------------------
    # HOJA 2: PARÁMETROS
    # ---------------------------------------------------------
    ws2 = wb.create_sheet(title="Parámetros")
    ws2.append(["Parámetro", "Valor Configurado"])
    for k, v in params.items():
        ws2.append([str(k), str(v)])

    # ---------------------------------------------------------
    # HOJA 3: RESUMEN
    # ---------------------------------------------------------
    ws3 = wb.create_sheet(title="Resumen Ejecución")
    ws3.append(["Métrica KPI", "Valor"])
    for k, v in kpis.items():
        ws3.append([str(k), str(v)])

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()